"""Generate the spreadsheet the domain reviewer fills in, pre-filled from content/*.json.

    python -m atlas_pipeline.review_template [output.xlsx]

No formulas on purpose: the file must behave the same in Excel, Numbers and Google Sheets, and the
JSON converter validates ids anyway. Yellow cells are the ones to edit.
"""

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import CONTENT_DIR, REPO_ROOT

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
INPUT_FILL = PatternFill("solid", fgColor="FFF9C4")
LOCKED_FILL = PatternFill("solid", fgColor="ECEFF1")
EXAMPLE_FILL = PatternFill("solid", fgColor="E3F2FD")
THIN = Side(style="thin", color="CFD8DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SIDE_KO = {"ipsi": "같은쪽", "contra": "반대쪽"}
CATEGORY_KO = {"somatosensory": "체성감각", "limbic": "둘레계통", "oculomotor": "안구운동·반사"}
REVIEW_CHOICES = ["맞음", "수정 필요", "모르겠음"]
SIDE_CHOICES = ["같은쪽", "반대쪽", "양쪽", "정중"]
CATEGORY_CHOICES = ["체성감각", "운동", "뇌신경 반사", "안구운동·반사", "시각", "청각·안뜰", "자율신경·시상하부",
                    "소뇌", "바닥핵", "둘레계통", "각성·수면", "통증 조절", "고위 겉질 기능", "척수 반사", "기타"]
PRIORITY_CHOICES = ["상", "중", "하"]
BLANK_ROWS = 60


def load_labels() -> dict[int, tuple[str, float]]:
    with open(REPO_ROOT / "docs" / "data" / "nextbrain-argmax-labels.tsv", newline="") as f:
        return {int(r["label_id"]): (r["name"], float(r["volume_mm3"])) for r in csv.DictReader(f, delimiter="\t")}


def structure_labels(source: dict, labels: dict) -> list[int]:
    ids = list(source.get("labels", []))
    for lo, hi in source.get("labelRanges", []):
        ids += [i for i in range(lo, hi + 1) if i in labels]
    return ids


def write_table(ws, headers, rows, *, widths, input_cols, notes=None, example_rows=0, blank_rows=BLANK_ROWS):
    """headers: list[str]; rows: list[list]; input_cols: header names the reviewer edits."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        if notes and h in notes:
            cell.comment = Comment(notes[h], "템플릿")
        ws.column_dimensions[get_column_letter(c)].width = widths.get(h, 16)
    ws.row_dimensions[1].height = 34
    total = len(rows) + blank_rows
    for r in range(total):
        values = rows[r] if r < len(rows) else [None] * len(headers)
        for c, (h, v) in enumerate(zip(headers, values), 1):
            cell = ws.cell(row=r + 2, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if r < example_rows:
                cell.fill = EXAMPLE_FILL
            elif h in input_cols:
                cell.fill = INPUT_FILL
            elif r < len(rows):
                cell.fill = LOCKED_FILL
            else:
                cell.fill = INPUT_FILL  # brand-new rows: every column is the reviewer's
    ws.freeze_panes = "C2" if len(headers) > 4 else "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{total + 1}"
    return total + 1


def add_dropdown(ws, header_names, header, choices, last_row):
    col = get_column_letter(header_names.index(header) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{last_row}")


def build(out_path: Path):
    structures = json.loads((CONTENT_DIR / "structures.json").read_text())
    pathways = json.loads((CONTENT_DIR / "pathways.json").read_text())["pathways"]
    labels = load_labels()
    groups = {gid: g["name"]["ko"] for gid, g in structures["groups"].items()}
    names = {s["id"]: s["name"]["ko"] for s in structures["structures"]}

    wb = Workbook()

    # ---- 안내 -------------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 118
    guide = [
        ("신경경로 3D 아틀라스 — 내용 검수·작성 시트", "title"),
        ("뷰어: https://neuro-pathway-atlas.vercel.app  (PC 브라우저 권장)", None),
        ("", None),
        ("어디를 고치면 되나요", "h"),
        ("· 노란 칸 = 채우거나 고칠 칸.   회색 칸 = 현재 시제품에 들어 있는 값(틀렸으면 '수정안' 칸에 적어 주세요).   파란 줄 = 작성 예시.", None),
        ("· 각 시트 아래쪽의 빈 노란 줄에 새 구조·경로·단계를 자유롭게 추가하면 됩니다. 확신이 없는 칸은 비워 두거나 '?' 로 두세요.", None),
        ("· 지금 들어 있는 경로 설명과 한국어 용어는 전부 AI가 쓴 초안입니다. 틀린 곳을 찾는 것이 1차 목표입니다.", None),
        ("", None),
        ("시트 구성", "h"),
        ("① 구조   뷰어에 나오는 3D 구조 한 개가 한 줄. '포함할_Allen_구조'에 여러 개를 적으면 하나의 메시로 합쳐집니다(묶기).", None),
        ("② 경로   경로 한 개가 한 줄 (이름·분류·요약·중요도).", None),
        ("③ 단계   경로의 단계 한 개가 한 줄. 가장 중요한 시트입니다.", None),
        ("④ 병변   경로별 병변 위치 → 증상.", None),
        ("⑤ 참고_Allen라벨   3D로 만들 수 있는 구조 전체 목록(333개, 읽기 전용). 이름은 Allen Human Brain Atlas의 구조명과 같습니다.", None),
        ("", None),
        ("꼭 지켜 주세요", "h"),
        ("1. 좌우는 '왼쪽/오른쪽'이 아니라 기시쪽 기준 '같은쪽/반대쪽'으로 적습니다. 뷰어에서 좌우 스위치를 누르면 전체가 뒤집히기 때문입니다.", None),
        ("   예) 뒤기둥-안쪽섬유띠 경로: 기시쪽 = 자극을 받는 몸 쪽 → 널판핵은 '같은쪽', 안쪽섬유띠·VPL·중심뒤이랑은 '반대쪽'.", None),
        ("2. 구조는 구조_id로 가리킵니다(① 시트의 A열). 한 단계에 여러 구조면 쉼표로 나열: bs-gracile, bs-cuneate", None),
        ("3. 3D에 아직 없는 구조도 지우지 말고 '메시_없는_구조' 칸에 적어 주세요. 다음에 무엇을 3D로 만들지 정하는 목록이 됩니다.", None),
        ("4. 새 구조를 추가할 때는 ⑤ 시트에서 이름을 찾아 '포함할_Allen_구조' 칸에 그대로 붙여 넣으면 됩니다. 목록에 없으면 이름만 적어 주세요(직접 모델링 후보).", None),
        ("5. 출처는 교과서 이름과 쪽수 정도면 충분합니다. 예) Blumenfeld 3판 p.279", None),
        ("", None),
        ("알려진 한계 (틀린 게 아닙니다)", "h"),
        ("· 오른쪽 반구는 왼쪽의 좌우 반전본.  · 뇌활은 가운데(몸통~다리) 구간의 데이터가 없어 경로선이 허공을 잇습니다.", None),
        ("· 눈돌림신경핵·혀밑신경핵은 데이터가 너무 작아 타원체로 대신했습니다('근사').  · 척수, 속섬유막, 말초 구조는 아직 없습니다.", None),
        ("· 부피 참고: 시상밑핵·가쪽무릎핵·둔덕이 문헌값보다 2~3.5배 크게 나옵니다. 눈으로 봐서 너무 커 보이는지 알려 주세요.", None),
    ]
    for r, (text, kind) in enumerate(guide, 1):
        cell = ws.cell(row=r, column=2, value=text)
        cell.font = Font(name=FONT, size=16 if kind == "title" else 11, bold=kind in ("title", "h"))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r, fill, label in [(len(guide) + 2, INPUT_FILL, "노란 칸: 채우거나 고치는 곳"), (len(guide) + 3, LOCKED_FILL, "회색 칸: 현재 값 (참고)"), (len(guide) + 4, EXAMPLE_FILL, "파란 줄: 작성 예시")]:
        c = ws.cell(row=r, column=2, value=label)
        c.fill, c.font, c.border = fill, Font(name=FONT, size=10), BORDER

    # ---- ① 구조 -----------------------------------------------------------------------------------
    ws = wb.create_sheet("구조")
    headers = ["구조_id", "한국어명", "영어명", "그룹", "포함할_Allen_구조", "부피_mm3(한쪽)", "형상", "검수", "수정안", "메모"]
    rows = []
    for s in structures["structures"]:
        src = s["source"]
        ids = structure_labels(src, labels) if src["atlas"] == "nextbrain" and "shell" not in src else []
        if ids:
            allen = ", ".join(labels[i][0] for i in ids)
            volume = round(sum(labels[i][1] for i in ids), 1)
        else:
            allen = "(뇌 윤곽: 자동 생성)" if "shell" in src else "(Harvard AAN 아틀라스)"
            volume = None
        shape = "근사(타원체)" if "placeholder" in src else "아틀라스"
        rows.append([s["id"], s["name"]["ko"], s["name"].get("en"), groups[s["group"]], allen, volume, shape, None, None, None])
    notes = {
        "구조_id": "영문 소문자와 하이픈. 새 구조는 비워 두면 개발자가 붙입니다.",
        "포함할_Allen_구조": "⑤ 참고_Allen라벨 시트의 이름을 쉼표로 나열. 여러 개를 적으면 한 덩어리로 합쳐집니다.",
        "부피_mm3(한쪽)": "아틀라스 라벨의 부피 합(정리 전 값). 문헌값과 비교하는 데 쓰세요.",
        "검수": "맞음 / 수정 필요 / 모르겠음",
        "수정안": "한국어 용어가 학교에서 쓰는 것과 다르거나, 묶음이 틀렸으면 여기에.",
    }
    last = write_table(ws, headers, rows, widths={"구조_id": 22, "한국어명": 30, "영어명": 34, "그룹": 16, "포함할_Allen_구조": 60, "부피_mm3(한쪽)": 13, "형상": 13, "검수": 12, "수정안": 36, "메모": 36},
                       input_cols={"검수", "수정안", "메모"}, notes=notes)
    add_dropdown(ws, headers, "검수", REVIEW_CHOICES, last)
    add_dropdown(ws, headers, "그룹", sorted(set(groups.values())), last)

    # ---- ② 경로 -----------------------------------------------------------------------------------
    ws = wb.create_sheet("경로")
    headers = ["경로_id", "한국어명", "영어명", "분류", "한줄요약", "기시쪽_설명", "중요도", "검수", "수정안·메모"]
    rows = [[p["id"], p["name"]["ko"], p["name"].get("en"), CATEGORY_KO.get(p["category"], p["category"]), p["summary"]["ko"], p["origin"]["ko"], None, None, None] for p in pathways]
    notes = {
        "경로_id": "영문 소문자. ③ 단계 · ④ 병변 시트에서 이 id로 연결됩니다.",
        "기시쪽_설명": "좌우 스위치 옆에 붙는 말. 예) 자극을 받는 몸 쪽 / 머리가 도는 쪽 / 빛을 비춘 눈",
        "중요도": "시험·수업에서의 중요도: 상 / 중 / 하",
    }
    last = write_table(ws, headers, rows, widths={"경로_id": 16, "한국어명": 30, "영어명": 38, "분류": 16, "한줄요약": 70, "기시쪽_설명": 22, "중요도": 9, "검수": 12, "수정안·메모": 44},
                       input_cols={"중요도", "검수", "수정안·메모"}, notes=notes, blank_rows=80)
    add_dropdown(ws, headers, "분류", CATEGORY_CHOICES, last)
    add_dropdown(ws, headers, "중요도", PRIORITY_CHOICES, last)
    add_dropdown(ws, headers, "검수", REVIEW_CHOICES, last)

    # ---- ③ 단계 -----------------------------------------------------------------------------------
    ws = wb.create_sheet("단계")
    headers = ["경로_id", "순서", "단계_제목", "뉴런_차수", "강조할_구조(구조_id)", "구조_이름(참고)", "좌우", "여기서_교차", "교차_위치",
               "신경전달물질", "설명", "메시_없는_구조", "출처", "검수", "수정안"]
    rows = []
    for p in pathways:
        for i, st in enumerate(p["steps"], 1):
            refs = st["structures"]
            sides = sorted({SIDE_KO[r["side"]] for r in refs})
            rows.append([
                p["id"], i, st["title"]["ko"], st.get("neuronOrder"),
                ", ".join(r["id"] for r in refs), ", ".join(names[r["id"]] for r in refs),
                " / ".join(sides) if sides else None, "예" if st.get("decussation") else None, None, None,
                st["description"]["ko"], (st.get("missing") or {}).get("ko"), None, None, None,
            ])
    notes = {
        "순서": "1, 2, 3 … 신호가 흐르는 순서.",
        "뉴런_차수": "1 / 2 / 3. 해당 없으면 빈칸.",
        "강조할_구조(구조_id)": "① 구조 시트의 구조_id. 여러 개면 쉼표로. 이 단계에서 밝게 표시됩니다.",
        "좌우": "기시쪽 기준: 같은쪽 / 반대쪽 / 양쪽 / 정중. 한 단계 안에서 구조마다 다르면 '수정안'에 풀어서 적어 주세요.",
        "여기서_교차": "이 단계에서 정중선을 건너면 '예'.",
        "교차_위치": "예) 숨뇌 아래쪽(감각교차), 앞백색맞교차, 위소뇌다리교차",
        "메시_없는_구조": "설명에는 나오지만 3D에 아직 없는 구조. 예) 속활꼴섬유, 속섬유막 뒤다리",
        "출처": "교과서와 쪽수. 예) Blumenfeld 3판 p.279",
    }
    last = write_table(ws, headers, rows, widths={"경로_id": 12, "순서": 6, "단계_제목": 30, "뉴런_차수": 8, "강조할_구조(구조_id)": 30, "구조_이름(참고)": 28, "좌우": 12, "여기서_교차": 10, "교차_위치": 24,
                                                 "신경전달물질": 14, "설명": 70, "메시_없는_구조": 28, "출처": 24, "검수": 12, "수정안": 44},
                       input_cols={"교차_위치", "신경전달물질", "출처", "검수", "수정안"}, notes=notes, blank_rows=150)
    add_dropdown(ws, headers, "좌우", SIDE_CHOICES, last)
    add_dropdown(ws, headers, "여기서_교차", ["예"], last)
    add_dropdown(ws, headers, "뉴런_차수", ["1", "2", "3", "4"], last)
    add_dropdown(ws, headers, "검수", REVIEW_CHOICES, last)

    # ---- ④ 병변 -----------------------------------------------------------------------------------
    ws = wb.create_sheet("병변")
    headers = ["경로_id", "병변_위치", "증상", "증상_쪽", "관련_증후군", "출처", "검수", "수정안"]
    rows = [[p["id"], l["level"]["ko"], l["deficit"]["ko"], None, None, None, None, None] for p in pathways for l in p["lesions"]]
    notes = {"증상_쪽": "병변 기준: 같은쪽 / 반대쪽 / 양쪽", "관련_증후군": "예) Brown-Séquard, Wallenberg, 핵사이눈근육마비(INO)"}
    last = write_table(ws, headers, rows, widths={"경로_id": 12, "병변_위치": 34, "증상": 70, "증상_쪽": 11, "관련_증후군": 26, "출처": 24, "검수": 12, "수정안": 44},
                       input_cols={"증상_쪽", "관련_증후군", "출처", "검수", "수정안"}, notes=notes, blank_rows=100)
    add_dropdown(ws, headers, "증상_쪽", ["같은쪽", "반대쪽", "양쪽"], last)
    add_dropdown(ws, headers, "검수", REVIEW_CHOICES, last)

    # ---- ⑤ 참고_Allen라벨 -------------------------------------------------------------------------
    ws = wb.create_sheet("참고_Allen라벨")
    used = {}
    for s in structures["structures"]:
        if s["source"]["atlas"] == "nextbrain" and "shell" not in s["source"]:
            for i in structure_labels(s["source"], labels):
                used[i] = s["id"]
    headers = ["라벨_id", "Allen_구조명", "부피_mm3(한쪽)", "크기_판정", "이미_쓰는_구조_id"]
    rows = []
    for i, (name, vol) in sorted(labels.items()):
        if i == 0:
            continue
        verdict = "메시 가능" if vol >= 20 else "작음(보정 필요)" if vol >= 5 else "극소(타원체로 대체)"
        rows.append([i, name, round(vol, 1), verdict, used.get(i)])
    write_table(ws, headers, rows, widths={"라벨_id": 9, "Allen_구조명": 62, "부피_mm3(한쪽)": 14, "크기_판정": 22, "이미_쓰는_구조_id": 24},
                input_cols=set(), blank_rows=0,
                notes={"Allen_구조명": "NextBrain 아틀라스(CC0)의 라벨. Allen Human Brain Atlas 온톨로지의 이름을 그대로 씁니다. ctx-rh-* 는 대뇌겉질 이랑(FreeSurfer 이름).",
                       "크기_판정": "아틀라스 볼륨에서의 부피 기준. 20mm3 이상이면 바로 3D로 만들 수 있고, 5 미만이면 위치만 쓰고 모양은 타원체로 대신합니다."})

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO_ROOT / "docs" / "handover-seongyeon" / "신경경로_내용검수_템플릿.xlsx"
    print(build(target))
